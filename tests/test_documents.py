from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from app.documents import DocumentError, extract_policy_text
from app.main import app
from app.report import PolicyReport, build_report_pdf


def test_extracts_utf8_text_file():
    assert extract_policy_text("policy.txt", b"\xef\xbb\xbfPrivacy policy\nWe keep your data.") == \
        "Privacy policy\nWe keep your data."


def test_extracts_docx_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Privacy policy")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Email"
    table.cell(0, 1).text = "Account access"
    output = BytesIO()
    document.save(output)

    text = extract_policy_text("policy.docx", output.getvalue())
    assert "Privacy policy" in text
    assert "Email | Account access" in text


def test_extracts_searchable_pdf_text():
    output = BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(72, 720, "Privacy policy and personal information")
    pdf.save()

    assert "Privacy policy and personal information" in extract_policy_text("policy.pdf", output.getvalue())


def test_rejects_scanned_pdf_without_selectable_text():
    output = BytesIO()
    canvas.Canvas(output).save()
    with pytest.raises(DocumentError, match="No selectable text"):
        extract_policy_text("scan.pdf", output.getvalue())


def test_builds_downloadable_pdf_report():
    report = PolicyReport.model_validate({
        "company": "Example Co",
        "overview": "A short overview.",
        "generated_at": "2026-10-23T10:00:00+00:00",
        "policy_sha256": "a" * 64,
        "model": "groq/test-model",
        "disclosures": [{
            "category": "data_sharing", "summary": "They share account data.",
            "evidence": "We share account data with providers.", "status": "disclosed",
            "evidence_verified": True,
        }],
        "stats": {"quotes_checked": 1, "quotes_verified": 1},
        "disclaimer": "Not legal advice.",
    })

    payload = build_report_pdf(report)
    extracted = " ".join((page.extract_text() or "") for page in PdfReader(BytesIO(payload)).pages)
    assert payload.startswith(b"%PDF-")
    assert "Example Co" in extracted
    assert "Verified quote" in extracted
    assert "Not legal advice." in extracted


def test_extract_endpoint_accepts_text_upload():
    response = TestClient(app).post(
        "/api/extract", files={"file": ("policy.txt", b"A sample privacy policy.", "text/plain")}
    )

    assert response.status_code == 200
    assert response.json()["policy_text"] == "A sample privacy policy."


def test_report_endpoint_returns_pdf():
    report = {
        "company": "Example Co", "overview": "A short overview.",
        "generated_at": "2026-10-23T10:00:00+00:00", "policy_sha256": "a" * 64,
        "model": "groq/test-model", "disclosures": [],
        "stats": {"quotes_checked": 0, "quotes_verified": 0}, "disclaimer": "Not legal advice.",
    }

    response = TestClient(app).post("/api/report", json=report)

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF-")